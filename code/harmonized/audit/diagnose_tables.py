"""Read-only manuscript diagnosis; all outputs are confined to this audit folder."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import sys
from zipfile import ZipFile

sys.dont_write_bytecode = True
from docx import Document
from lxml import etree
import pdfplumber
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
SOURCE = PROJECT / 'manuscript/staging_v15/IJISMD_v15_editorial_working_draft.docx'
PDF = PROJECT / 'manuscript/qa_projection/IJISMD_v15_editorial_working_draft.pdf'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(element):
    return ''.join(element.xpath('.//w:t/text()', namespaces=NS))


def main():
    snapshot = ROOT / 'snapshot'
    snapshot.mkdir(exist_ok=True)
    for path in (SOURCE, PDF):
        dest = snapshot / path.name
        if not dest.exists():
            shutil.copy2(path, dest)
    source = snapshot / SOURCE.name
    pdf_path = snapshot / PDF.name
    with ZipFile(source) as archive:
        xml = etree.fromstring(archive.read('word/document.xml'))
        styles = etree.fromstring(archive.read('word/styles.xml'))
    body = xml.find('w:body', NS)
    tables = body.findall('w:tbl', NS)
    table_info = []
    for i, table in enumerate(tables):
        index = body.index(table)
        item = {'table_number': i + 1, 'body_index': index,
                'preceding': [text(p) for p in list(body)[max(0, index - 3):index]],
                'following': [text(p) for p in list(body)[index + 1:index + 4]],
                'tblPr': etree.tostring(table.find('w:tblPr', NS), encoding='unicode'),
                'rows': []}
        for ri, row in enumerate(table.findall('w:tr', NS)):
            item['rows'].append({'index': ri, 'text': [text(c) for c in row.findall('w:tc', NS)],
                                 'xml': etree.tostring(row, encoding='unicode')})
        table_info.append(item)
        if i in (5, 7):
            (ROOT / f'table_{i+1}.xml').write_bytes(etree.tostring(table, pretty_print=True))
    (ROOT / 'styles.xml').write_bytes(etree.tostring(styles, pretty_print=True))
    doc = Document(source)
    pages = []
    with pdfplumber.open(pdf_path) as pdf:
        for index, page in enumerate(pdf.pages):
            value = page.extract_text() or ''
            pages.append(value)
            if 'Mapping controls' in value or '37 predictor' in value or index == 11:
                with pdfium.PdfDocument(pdf_path) as rendered:
                    rendered[index].render(scale=1.25).to_pil().save(ROOT / f'original_page_{index+1}.png')
    all_text = '\n'.join(pages)
    (ROOT / 'original_pdf_text.txt').write_text(all_text, encoding='utf-8')
    report = {'source': str(SOURCE), 'pdf': str(PDF),
              'source_sha256_snapshot': sha(source), 'source_sha256_current': sha(SOURCE),
              'pdf_sha256_snapshot': sha(pdf_path), 'pdf_sha256_current': sha(PDF),
              'pdf_pages': len(pages), 'tables': table_info,
              'pdf_needles': {needle: [i+1 for i, p in enumerate(pages) if needle in p]
                              for needle in ('Table 6', 'Random pseudo maps', '0.8958', '0.8814', 'Table 8')},
              'python_docx_table_counts': [len(t.rows) for t in doc.tables]}
    (ROOT / 'structural_diagnosis.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps({key: value for key, value in report.items() if key != 'tables'}, indent=2))
    for i in (5, 7):
        table = table_info[i]
        print('TABLE', i + 1, 'BEFORE', table['preceding'], 'AFTER', table['following'])
        for row in table['rows']:
            print(row['index'], row['text'])
        print('PROPERTIES', table['tblPr'])


if __name__ == '__main__':
    main()
