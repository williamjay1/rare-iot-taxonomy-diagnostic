"""Independent cell-level numeric and rendered-table audits on the saved snapshot."""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import statistics
import pypdfium2 as pdfium

from docx import Document
import pdfplumber

ROOT = Path(__file__).resolve().parent
WORK = ROOT.parents[1]
SOURCE = ROOT / 'snapshot/IJISMD_v15_editorial_working_draft.docx'
PDF = ROOT / 'snapshot/IJISMD_v15_editorial_working_draft.pdf'

METHODS = {
    'Flat XGBoost': 'flat_weighted_xgboost',
    'Flat LightGBM': 'flat_weighted_lightgbm',
    'Flat equal blend': 'flat_equal_probability_blend',
    'Calibrated flat blend': 'flat_temperature_calibrated_blend',
    'Flat stacker': 'flat_newton_log_probability_stacker',
    'Soft real hierarchy': 'soft_taxonomy_factorisation',
    'Hard real hierarchy': 'hard_top_down_taxonomy',
    'Real blend': 'equal_flat_taxonomy_blend',
    'Cyclic blend': 'equal_flat_cyclic_blend',
    'Support blend': 'equal_flat_support_matched_blend',
}


def norm(value):
    return ' '.join((value or '').split())


def main():
    handoff = json.loads((WORK / 'harmonized/MANUSCRIPT_HANDOFF.json').read_text())
    assert handoff['completion'] is True
    doc = Document(SOURCE)
    table = doc.tables[7]
    checks = []
    protocol_rows = {'paired46': 0, 'grouped37': 0}
    for index, row in enumerate(table.rows[1:], 1):
        cells = [cell.text for cell in row.cells]
        protocol = {'Original roles': 'paired46', 'Regrouped': 'grouped37'}[cells[0]]
        protocol_rows[protocol] += 1
        method = METHODS[cells[1]]
        contract = handoff['protocols'][protocol]
        with Path(contract['result_csv']).open(newline='') as handle:
            source = [r for r in csv.DictReader(handle) if r['run_id'] == method]
        assert sorted(int(r['seed']) for r in source) == [42, 43, 44, 45, 46, 47]
        ranges = {(r['run_id'], r['metric']): r for r in contract['ranges']}
        for column, metric in enumerate(('f1_macro', 'rare_recall', 'ece_10'), 2):
            values = [float(r[metric]) for r in source]
            mean, sd = statistics.mean(values), statistics.stdev(values)
            frozen = ranges[method, metric]
            assert abs(mean-frozen['mean']) < 1e-14
            assert abs(sd-frozen['sample_sd']) < 1e-14
            expected = f'{mean:.4f} ({sd:.4f})'
            checks.append({'row_index': index, 'protocol': protocol, 'run_id': method,
                           'metric': metric, 'actual': cells[column], 'expected': expected,
                           'pass': cells[column] == expected, 'csv': contract['result_csv'],
                           'mean_full_precision': mean, 'sample_sd_full_precision': sd})
    pdf_tables = []
    with pdfplumber.open(PDF) as pdf:
        for index, page in enumerate(pdf.pages):
            for pdf_table in page.extract_tables():
                pdf_tables.append({'page': index+1, 'cells': [[norm(c) for c in row] for row in pdf_table]})
    pdf_checks = {}
    for number in (6, 8, 14):
        expected_rows = [[norm(c.text) for c in r.cells] for r in doc.tables[number-1].rows]
        header = expected_rows[0]
        matched_tables = [t for t in pdf_tables if t['cells'] and t['cells'][0] == header]
        flat_rows = [(t['page'], row) for t in matched_tables for row in t['cells'][1:]]
        matches = [{'source_row': i, 'cells': row,
                    'pages': [page for page, actual in flat_rows
                              if [''.join(c.split()) for c in actual] == [''.join(c.split()) for c in row]],
                    'exact_word_spacing_pages': [page for page, actual in flat_rows if actual == row]}
                   for i, row in enumerate(expected_rows[1:], 1)]
        pdf_checks[str(number)] = {'source_data_rows': len(expected_rows)-1,
                                   'matched_pdf_tables': matched_tables,
                                   'row_matches': matches,
                                   'complete': all(item['pages'] for item in matches)}
    report = {'source_docx': str(SOURCE), 'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
              'pdf': str(PDF), 'table8_total_rows_including_header': len(table.rows),
              'table8_protocol_data_rows': protocol_rows,
              'table8_numeric_cells_checked': len(checks),
              'table8_numeric_values_checked': len(checks)*2,
              'table8_numeric_status': 'PASS' if all(c['pass'] for c in checks) else 'FAIL',
              'checks': checks, 'pdf_cell_audit': pdf_checks}
    (ROOT / 'projection_and_pdf_cell_audit.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    with pdfium.PdfDocument(PDF) as pdf:
        pdf[23].render(scale=1.25).to_pil().save(ROOT / 'original_page_24.png')
    print(json.dumps({k:v for k,v in report.items() if k not in ('checks', 'pdf_cell_audit')}, indent=2))
    for number, audit in pdf_checks.items():
        print('TABLE', number, json.dumps(audit, indent=2))


if __name__ == '__main__':
    main()
